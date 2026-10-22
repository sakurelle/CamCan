from __future__ import annotations

import csv
import json
import os
import sqlite3
import tempfile
import uuid
import shutil
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterator

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from marker_qc.config import Norms
from marker_qc.models import Decision, InspectionDraft


SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS batches (
    id INTEGER PRIMARY KEY,
    code TEXT NOT NULL UNIQUE,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS markers (
    id INTEGER PRIMARY KEY,
    batch_id INTEGER NOT NULL REFERENCES batches(id),
    code TEXT NOT NULL,
    created_at TEXT NOT NULL,
    UNIQUE(batch_id, code)
);
CREATE TABLE IF NOT EXISTS inspections (
    id TEXT PRIMARY KEY,
    marker_id INTEGER NOT NULL REFERENCES markers(id),
    stage TEXT NOT NULL,
    created_at TEXT NOT NULL,
    operator_name TEXT NOT NULL,
    length_mm REAL NOT NULL,
    width_mm REAL NOT NULL,
    length_ok INTEGER NOT NULL,
    width_ok INTEGER NOT NULL,
    dimensions_ok INTEGER NOT NULL,
    ends_ok INTEGER NOT NULL,
    shape_ok INTEGER NOT NULL,
    decision TEXT NOT NULL,
    automatic INTEGER NOT NULL,
    confidence REAL NOT NULL,
    scale_mm_per_px REAL NOT NULL,
    points_json TEXT NOT NULL,
    diagnostics_json TEXT NOT NULL,
    norms_json TEXT NOT NULL,
    note TEXT NOT NULL,
    raw_path TEXT NOT NULL,
    roi_path TEXT NOT NULL,
    mask_path TEXT NOT NULL,
    overlay_path TEXT NOT NULL,
    UNIQUE(marker_id, stage)
);
CREATE TABLE IF NOT EXISTS audit_log (
    id INTEGER PRIMARY KEY,
    created_at TEXT NOT NULL,
    operator_name TEXT NOT NULL,
    event_type TEXT NOT NULL,
    details_json TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_inspections_created ON inspections(created_at);
"""


class StorageError(RuntimeError):
    pass


@dataclass(frozen=True)
class ClearHistoryResult:
    deleted_records: int
    deleted_files: int
    warnings: tuple[str, ...] = ()


class Storage:
    def __init__(self, database: Path, images_dir: Path):
        self.database = Path(database)
        self.images_dir = Path(images_dir)
        self.database.parent.mkdir(parents=True, exist_ok=True)
        self.images_dir.mkdir(parents=True, exist_ok=True)
        self._initialize()

    @contextmanager
    def connect(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.database, timeout=15)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = WAL")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def _initialize(self) -> None:
        with self.connect() as connection:
            connection.executescript(SCHEMA)

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")

    def audit(self, operator: str, event_type: str, details: dict) -> None:
        with self.connect() as connection:
            connection.execute(
                "INSERT INTO audit_log(created_at,operator_name,event_type,details_json) VALUES(?,?,?,?)",
                (self._now(), operator, event_type, json.dumps(details, ensure_ascii=False)),
            )

    def marker_decisions(self, batch_code: str, marker_code: str) -> list[Decision]:
        with self.connect() as connection:
            rows = connection.execute(
                """SELECT i.decision FROM inspections i
                   JOIN markers m ON m.id=i.marker_id JOIN batches b ON b.id=m.batch_id
                   WHERE b.code=? AND m.code=? ORDER BY i.created_at""",
                (batch_code, marker_code),
            ).fetchall()
        return [Decision(row[0]) for row in rows]

    def _batch_and_marker(self, connection: sqlite3.Connection, batch: str, marker: str) -> int:
        now = self._now()
        connection.execute("INSERT OR IGNORE INTO batches(code,created_at) VALUES(?,?)", (batch, now))
        batch_id = int(connection.execute("SELECT id FROM batches WHERE code=?", (batch,)).fetchone()[0])
        connection.execute(
            "INSERT OR IGNORE INTO markers(batch_id,code,created_at) VALUES(?,?,?)", (batch_id, marker, now)
        )
        return int(connection.execute(
            "SELECT id FROM markers WHERE batch_id=? AND code=?", (batch_id, marker)
        ).fetchone()[0])

    def _write_image(self, directory: Path, name: str, image: np.ndarray) -> Path:
        directory.mkdir(parents=True, exist_ok=True)
        success, encoded = cv2.imencode(".png", image)
        if not success:
            raise StorageError(f"Не удалось подготовить изображение {name} к записи.")
        handle, temporary = tempfile.mkstemp(prefix=name + "_", suffix=".tmp", dir=directory)
        try:
            with os.fdopen(handle, "wb") as stream:
                stream.write(encoded.tobytes())
                stream.flush()
                os.fsync(stream.fileno())
            final = directory / f"{name}.png"
            os.replace(temporary, final)
            return final
        except Exception:
            try:
                os.unlink(temporary)
            except OSError:
                pass
            raise

    def save_inspection(
        self,
        draft: InspectionDraft,
        norms: Norms,
        images: dict[str, np.ndarray],
    ) -> str:
        required = {"raw", "roi", "mask", "overlay"}
        if required != set(images):
            raise StorageError("Для записи нужны изображения raw, roi, mask и overlay.")
        inspection_id = uuid.uuid4().hex
        directory = self.images_dir / datetime.now().strftime("%Y") / datetime.now().strftime("%m") / inspection_id
        saved = {kind: self._write_image(directory, kind, image) for kind, image in images.items()}
        result = draft.result
        try:
            with self.connect() as connection:
                marker_id = self._batch_and_marker(connection, draft.batch_code, draft.marker_code)
                connection.execute(
                    """INSERT INTO inspections(
                       id,marker_id,stage,created_at,operator_name,length_mm,width_mm,length_ok,width_ok,
                       dimensions_ok,ends_ok,shape_ok,decision,automatic,confidence,scale_mm_per_px,
                       points_json,diagnostics_json,norms_json,note,raw_path,roi_path,mask_path,overlay_path
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        inspection_id, marker_id, draft.stage.value, self._now(), draft.operator_name,
                        result.length_mm, result.width_mm, int(result.length_ok), int(result.width_ok),
                        int(result.dimensions_ok), int(draft.ends_ok), int(draft.shape_ok), draft.decision.value,
                        int(result.automatic), result.confidence, result.scale_mm_per_px,
                        json.dumps(result.points.as_json(), ensure_ascii=False),
                        json.dumps(result.diagnostics, ensure_ascii=False),
                        json.dumps(norms.__dict__, ensure_ascii=False), draft.note,
                        *(str(saved[key]) for key in ("raw", "roi", "mask", "overlay")),
                    ),
                )
        except sqlite3.IntegrityError as exc:
            raise StorageError("Этот этап ОТК для маркера уже сохранён.") from exc
        except sqlite3.Error as exc:
            raise StorageError(f"Ошибка базы данных: {exc}") from exc
        return inspection_id

    def history(self, limit: int = 1000) -> list[sqlite3.Row]:
        with self.connect() as connection:
            return connection.execute(
                """SELECT i.created_at,b.code AS batch_code,m.code AS marker_code,i.stage,i.operator_name,
                          i.length_mm,i.width_mm,i.length_ok,i.width_ok,i.dimensions_ok,i.ends_ok,i.shape_ok,i.decision,
                          i.automatic,i.confidence,i.scale_mm_per_px,i.points_json,i.overlay_path
                   FROM inspections i JOIN markers m ON m.id=i.marker_id JOIN batches b ON b.id=m.batch_id
                   ORDER BY i.created_at DESC LIMIT ?""",
                (int(limit),),
            ).fetchall()

    def batch_counts(self, batch_code: str) -> dict[str, int]:
        with self.connect() as connection:
            rows = connection.execute(
                """SELECT i.decision,COUNT(*) FROM inspections i
                   JOIN markers m ON m.id=i.marker_id JOIN batches b ON b.id=m.batch_id
                   WHERE b.code=? GROUP BY i.decision""",
                (batch_code,),
            ).fetchall()
        result = {decision.value: 0 for decision in Decision}
        result.update({str(row[0]): int(row[1]) for row in rows})
        return result

    def export_csv(self, destination: Path) -> None:
        rows = self.history(limit=1_000_000)
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        with destination.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow([
                "Дата", "Партия", "Маркер", "Этап", "Оператор", "Длина_мм", "Ширина_мм",
                "Длина_годна", "Ширина_годна", "Размеры_годны", "Торцы_ровные", "Форма_однородная",
                "Решение", "Автоматические_точки", "Уверенность", "Масштаб_мм_на_px", "Точки_JSON", "Изображение",
            ])
            for row in reversed(rows):
                writer.writerow(list(row))

    def export_xlsx(self, destination: Path) -> None:
        rows = self.history(limit=1_000_000)
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Измерения"
        headers = ("Дата и время", "Длина, мм", "Ширина, мм", "Итог", "Точки")
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        approved_fill = PatternFill("solid", fgColor="C6EFCE")
        rejected_fill = PatternFill("solid", fgColor="FFC7CE")
        for row in rows:
            created = datetime.fromisoformat(str(row["created_at"]))
            if created.tzinfo is not None:
                created = created.replace(tzinfo=None)
            excel_row = sheet.max_row + 1
            sheet.append((
                created,
                float(row["length_mm"]),
                float(row["width_mm"]),
                "ДА" if row["dimensions_ok"] else "НЕТ",
                "Авто" if row["automatic"] else "Вручную",
            ))
            sheet.cell(excel_row, 1).number_format = "dd.mm.yyyy hh:mm:ss"
            sheet.cell(excel_row, 2).number_format = "0.0000"
            sheet.cell(excel_row, 3).number_format = "0.0000"
            sheet.cell(excel_row, 4).fill = approved_fill if row["dimensions_ok"] else rejected_fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:E{max(sheet.max_row, 1)}"
        for column, width in {"A": 21, "B": 14, "C": 14, "D": 11, "E": 13}.items():
            sheet.column_dimensions[column].width = width
        try:
            workbook.save(destination)
        except Exception as exc:
            raise StorageError(f"Ошибка создания Excel-файла: {exc}") from exc

    def clear_history(self) -> ClearHistoryResult:
        """Удалить измерения и только принадлежащие им файлы внутри images_dir."""
        with self.connect() as connection:
            rows = connection.execute(
                "SELECT raw_path,roi_path,mask_path,overlay_path FROM inspections"
            ).fetchall()
        if not rows:
            return ClearHistoryResult(0, 0)

        images_root = self.images_dir.resolve()
        managed_files: list[Path] = []
        warnings: list[str] = []
        seen: set[Path] = set()
        for row in rows:
            for raw_path in row:
                candidate = Path(str(raw_path))
                if not candidate.is_absolute():
                    warnings.append(f"Пропущен небезопасный относительный путь: {candidate}")
                    continue
                resolved = candidate.resolve(strict=False)
                try:
                    resolved.relative_to(images_root)
                except ValueError:
                    warnings.append(f"Файл вне каталога снимков не удалён: {candidate}")
                    continue
                if resolved not in seen and resolved.is_file():
                    seen.add(resolved)
                    managed_files.append(resolved)

        quarantine = images_root / ".history-trash" / uuid.uuid4().hex
        copied: list[tuple[Path, Path]] = []
        try:
            for source in managed_files:
                relative = source.relative_to(images_root)
                backup = quarantine / relative
                backup.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, backup)
                copied.append((source, backup))
        except OSError as exc:
            shutil.rmtree(quarantine, ignore_errors=True)
            raise StorageError(f"Очистка отменена: не удалось подготовить резервную копию снимков: {exc}") from exc

        try:
            with self.connect() as connection:
                deleted_records = int(connection.execute("SELECT COUNT(*) FROM inspections").fetchone()[0])
                connection.execute("DELETE FROM inspections")
                connection.execute("DELETE FROM markers WHERE NOT EXISTS (SELECT 1 FROM inspections WHERE marker_id=markers.id)")
                connection.execute("DELETE FROM batches WHERE NOT EXISTS (SELECT 1 FROM markers WHERE batch_id=batches.id)")
                connection.execute(
                    "INSERT INTO audit_log(created_at,operator_name,event_type,details_json) VALUES(?,?,?,?)",
                    (self._now(), "", "HISTORY_CLEARED", json.dumps({"records": deleted_records}, ensure_ascii=False)),
                )
        except sqlite3.Error as exc:
            shutil.rmtree(quarantine, ignore_errors=True)
            raise StorageError(f"Очистка отменена из-за ошибки базы данных: {exc}") from exc

        deleted_files = 0
        file_errors = False
        for source, _backup in copied:
            try:
                source.unlink()
                deleted_files += 1
            except OSError as exc:
                file_errors = True
                warnings.append(f"Не удалось удалить снимок {source.name}: {exc}")
        if file_errors:
            warnings.append(f"Резервные копии сохранены для диагностики: {quarantine}")
        else:
            try:
                shutil.rmtree(quarantine)
            except OSError as exc:
                warnings.append(f"История удалена, но временный каталог не очищен: {exc}")
        return ClearHistoryResult(deleted_records, deleted_files, tuple(warnings))

