import csv
import numbers
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from marker_qc.config import Norms
from marker_qc.models import (
    Decision, InspectionDraft, InspectionStage, MeasurementPoints, MeasurementResult,
)
from marker_qc.storage import Storage, StorageError


def draft(stage=InspectionStage.FIRST, decision=Decision.APPROVED_STAGE1, marker="М-001", automatic=True):
    points = MeasurementPoints((0, 0), (350, 0), (0, 0), (350, 0), (175, 58))
    result = MeasurementResult(7.0, 1.16, True, True, True, points, 0.02, automatic, 0.9)
    return InspectionDraft("Партия А", marker, "Оператор", stage, result, True, True, decision, "норма")


def test_storage_survives_reopen_and_exports_utf8_bom(tmp_path) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    image = np.zeros((20, 20, 3), dtype=np.uint8)
    inspection_id = storage.save_inspection(draft(), Norms(), {key: image for key in ("raw", "roi", "mask", "overlay")})
    assert len(inspection_id) == 32
    reopened = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    assert reopened.marker_decisions("Партия А", "М-001") == [Decision.APPROVED_STAGE1]
    destination = tmp_path / "экспорт.csv"
    reopened.export_csv(destination)
    assert destination.read_bytes().startswith(b"\xef\xbb\xbf")
    with destination.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream, delimiter=";"))
    assert rows[1][1:3] == ["Партия А", "М-001"]


def test_duplicate_stage_is_prevented(tmp_path) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    image = np.zeros((10, 10, 3), dtype=np.uint8)
    images = {key: image for key in ("raw", "roi", "mask", "overlay")}
    storage.save_inspection(draft(), Norms(), images)
    with pytest.raises(StorageError, match="уже сохранён"):
        storage.save_inspection(draft(), Norms(), images)


def test_clear_history_removes_records_and_managed_images_only(tmp_path) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    image = np.zeros((10, 10, 3), dtype=np.uint8)
    images = {key: image for key in ("raw", "roi", "mask", "overlay")}
    storage.save_inspection(draft(marker="M-1"), Norms(), images)
    storage.save_inspection(draft(marker="M-2"), Norms(), images)
    unrelated = storage.images_dir / "keep.txt"
    unrelated.write_text("keep", encoding="utf-8")
    result = storage.clear_history()
    assert result.deleted_records == 2
    assert result.deleted_files == 8
    assert storage.history() == []
    assert unrelated.read_text(encoding="utf-8") == "keep"
    assert storage.clear_history().deleted_records == 0
    storage.save_inspection(draft(marker="M-3"), Norms(), images)
    assert len(storage.history()) == 1


def test_clear_history_never_deletes_path_outside_images(tmp_path) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    image = np.zeros((10, 10, 3), dtype=np.uint8)
    storage.save_inspection(draft(), Norms(), {key: image for key in ("raw", "roi", "mask", "overlay")})
    outside = tmp_path / "outside.png"
    outside.write_bytes(b"not managed")
    with storage.connect() as connection:
        connection.execute("UPDATE inspections SET raw_path=?", (str(outside),))
    result = storage.clear_history()
    assert outside.exists()
    assert any("вне каталога" in warning for warning in result.warnings)


def test_export_xlsx_has_typed_values_and_formatting(tmp_path) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    image = np.zeros((10, 10, 3), dtype=np.uint8)
    storage.save_inspection(draft(automatic=False), Norms(), {key: image for key in ("raw", "roi", "mask", "overlay")})
    before = len(storage.history())
    destination = tmp_path / "измерения.xlsx"
    storage.export_xlsx(destination)
    workbook = load_workbook(destination)
    sheet = workbook["Измерения"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:E2"
    assert isinstance(sheet["A2"].value, __import__("datetime").datetime)
    assert isinstance(sheet["B2"].value, numbers.Real)
    assert sheet["B2"].number_format == "0.0000"
    assert sheet["D2"].value == "ДА"
    assert sheet["E2"].value == "Вручную"
    assert len(storage.history()) == before


def test_export_empty_xlsx_is_valid(tmp_path) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    destination = tmp_path / "empty.xlsx"
    storage.export_xlsx(destination)
    sheet = load_workbook(destination)["Измерения"]
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == ["Дата и время", "Длина, мм", "Ширина, мм", "Итог", "Точки"]


def test_clear_history_reports_file_deletion_failure(tmp_path, monkeypatch) -> None:
    storage = Storage(tmp_path / "qc.sqlite3", tmp_path / "images")
    image = np.zeros((10, 10, 3), dtype=np.uint8)
    storage.save_inspection(draft(), Norms(), {key: image for key in ("raw", "roi", "mask", "overlay")})
    original_unlink = Path.unlink

    def fail_one_file(path, *args, **kwargs):
        if path.name == "raw.png" and ".history-trash" not in path.parts:
            raise OSError("simulated access error")
        return original_unlink(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", fail_one_file)
    result = storage.clear_history()
    assert storage.history() == []
    assert any("simulated access error" in warning for warning in result.warnings)
    assert any("Резервные копии" in warning for warning in result.warnings)

